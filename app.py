from flask import Flask, render_template, request, send_from_directory
from pypdf import PdfReader
from docx import Document
from werkzeug.utils import secure_filename

import os
import hashlib


# ============================================================
# ReportLab compatibility fix
# ============================================================

_original_md5 = hashlib.md5


def compatible_md5(data=b"", *args, **kwargs):
    """
    Compatibility wrapper for environments where
    openssl_md5() does not accept usedforsecurity.
    """
    return _original_md5(data)


hashlib.md5 = compatible_md5


from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle
)


# ============================================================
# Flask Configuration
# ============================================================

app = Flask(__name__)

UPLOAD_FOLDER = "uploads"
REPORT_FOLDER = "reports"

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
app.config["REPORT_FOLDER"] = REPORT_FOLDER

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(REPORT_FOLDER, exist_ok=True)


# ============================================================
# Skills
# ============================================================

SKILLS = [
    "python",
    "java",
    "aws",
    "linux",
    "docker",
    "kubernetes",
    "terraform",
    "jenkins",
    "git",
    "github",
    "sql",
    "mysql",
    "flask",
    "html",
    "css",
    "javascript",
    "machine learning",
    "devops",
    "cloud computing"
]


# ============================================================
# Skill Recommendations
# ============================================================

SKILL_RECOMMENDATIONS = {

    "python":
        "Improve Python programming and automation skills.",

    "java":
        "Learn Java fundamentals and object-oriented programming.",

    "aws":
        "Learn AWS core services such as EC2, S3, IAM, VPC and CloudWatch.",

    "linux":
        "Practice Linux commands, system administration and troubleshooting.",

    "docker":
        "Learn Docker images, containers, Dockerfiles and Docker Compose.",

    "kubernetes":
        "Learn Kubernetes pods, deployments, services and basic cluster management.",

    "terraform":
        "Learn Terraform infrastructure as code and AWS resource provisioning.",

    "jenkins":
        "Learn Jenkins pipelines and CI/CD automation.",

    "git":
        "Practice Git commands, branching, merging and version control.",

    "github":
        "Learn GitHub repositories, pull requests and collaboration workflows.",

    "sql":
        "Practice SQL queries, joins, filtering and database operations.",

    "mysql":
        "Learn MySQL databases, tables, queries and database management.",

    "flask":
        "Learn Flask routes, templates, forms and REST APIs.",

    "html":
        "Improve HTML structure, forms and semantic elements.",

    "css":
        "Improve CSS layouts, responsive design and styling.",

    "javascript":
        "Learn JavaScript fundamentals and browser interactions.",

    "machine learning":
        "Learn machine learning fundamentals and common algorithms.",

    "devops":
        "Learn CI/CD, Docker, cloud infrastructure and monitoring.",

    "cloud computing":
        "Learn cloud computing concepts and AWS fundamentals."
}


# ============================================================
# PDF Text Extraction
# ============================================================

def extract_pdf_text(file_path):

    reader = PdfReader(file_path)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:

            text += page_text.lower() + "\n"

    return text


# ============================================================
# DOCX Text Extraction
# ============================================================

def extract_docx_text(file_path):

    document = Document(file_path)

    text = ""

    # Normal paragraphs
    for paragraph in document.paragraphs:

        if paragraph.text:

            text += paragraph.text.lower() + "\n"

    # Tables
    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text:

                    text += cell.text.lower() + "\n"

    return text


# ============================================================
# Resume Text Extraction
# ============================================================

def extract_resume_text(file_path):

    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":

        return extract_pdf_text(file_path)

    elif extension == ".docx":

        return extract_docx_text(file_path)

    else:

        raise ValueError("Unsupported file format")


# ============================================================
# Find Skills
# ============================================================

def find_skills(text):

    found_skills = []

    text = text.lower()

    for skill in SKILLS:

        if skill.lower() in text:

            found_skills.append(skill)

    return found_skills


# ============================================================
# Recommendations
# ============================================================

def generate_recommendations(missing_skills):

    recommendations = []

    for skill in missing_skills:

        recommendation = SKILL_RECOMMENDATIONS.get(skill)

        if recommendation:

            recommendations.append({

                "skill": skill,

                "recommendation": recommendation

            })

    return recommendations


# ============================================================
# Generate PDF Report
# ============================================================

def generate_pdf_report(
    filename,
    match_percentage,
    matched_skills,
    missing_skills,
    recommendations
):

    report_path = os.path.join(
        app.config["REPORT_FOLDER"],
        filename
    )

    document = SimpleDocTemplate(
        report_path,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=22,
        spaceAfter=20
    )

    heading_style = ParagraphStyle(
        "Heading",
        parent=styles["Heading2"],
        fontSize=15,
        spaceBefore=15,
        spaceAfter=10
    )

    normal_style = ParagraphStyle(
        "NormalCustom",
        parent=styles["Normal"],
        fontSize=10,
        leading=15
    )

    story = []

    # Title
    story.append(
        Paragraph(
            "AWS Job Detector - Analysis Report",
            title_style
        )
    )

    story.append(Spacer(1, 10))

    # Match Score
    story.append(
        Paragraph(
            f"<b>Resume Match Score:</b> {match_percentage}%",
            heading_style
        )
    )

    story.append(Spacer(1, 10))

    # Matched Skills
    story.append(
        Paragraph(
            "Matched Skills",
            heading_style
        )
    )

    if matched_skills:

        matched_text = ", ".join(
            skill.title()
            for skill in matched_skills
        )

    else:

        matched_text = "No matching skills found."

    story.append(
        Paragraph(
            matched_text,
            normal_style
        )
    )

    # Missing Skills
    story.append(
        Paragraph(
            "Missing Skills",
            heading_style
        )
    )

    if missing_skills:

        missing_text = ", ".join(
            skill.title()
            for skill in missing_skills
        )

    else:

        missing_text = "No major missing skills found."

    story.append(
        Paragraph(
            missing_text,
            normal_style
        )
    )

    # Recommendations
    story.append(
        Paragraph(
            "Recommended Skills to Learn",
            heading_style
        )
    )

    if recommendations:

        recommendation_data = [
            ["Skill", "Recommendation"]
        ]

        for item in recommendations:

            recommendation_data.append([
                item["skill"].title(),
                item["recommendation"]
            ])

        table = Table(
            recommendation_data,
            colWidths=[120, 350]
        )

        table.setStyle(
            TableStyle([
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.black
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP"
                ),
                (
                    "FONTNAME",
                    (0, 1),
                    (-1, -1),
                    "Helvetica"
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    9
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    8
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    8
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    8
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    8
                )
            ])
        )

        story.append(table)

    else:

        story.append(
            Paragraph(
                "No additional skills recommendations.",
                normal_style
            )
        )

    story.append(Spacer(1, 30))

    story.append(
        Paragraph(
            "Generated by AWS Job Detector",
            normal_style
        )
    )

    document.build(story)

    return filename


# ============================================================
# Home Route
# ============================================================

@app.route("/", methods=["GET", "POST"])
def index():

    result = None
    error = None

    if request.method == "POST":

        # Get resume
        resume = request.files.get("resume")

        # Get job description
        job_description = request.form.get(
            "job_description",
            ""
        )

        # Check resume
        if not resume or resume.filename == "":

            error = "Please upload a PDF or DOCX resume."

            return render_template(
                "index.html",
                error=error
            )

        # Secure filename
        filename = secure_filename(
            resume.filename
        )

        # Get extension
        extension = os.path.splitext(
            filename
        )[1].lower()

        # Allowed formats
        allowed_extensions = [
            ".pdf",
            ".docx"
        ]

        if extension not in allowed_extensions:

            error = "Please upload a PDF or DOCX file only."

            return render_template(
                "index.html",
                error=error
            )

        # Save resume
        resume_path = os.path.join(
            app.config["UPLOAD_FOLDER"],
            filename
        )

        resume.save(resume_path)

        try:

            # Extract resume text
            resume_text = extract_resume_text(
                resume_path
            )

            # Resume skills
            resume_skills = find_skills(
                resume_text
            )

            # Job skills
            job_skills = find_skills(
                job_description
            )

            # Matched skills
            matched_skills = sorted(
                list(
                    set(resume_skills)
                    &
                    set(job_skills)
                )
            )

            # Missing skills
            missing_skills = sorted(
                list(
                    set(job_skills)
                    -
                    set(resume_skills)
                )
            )

            # Match percentage
            if job_skills:

                match_percentage = round(
                    len(matched_skills)
                    /
                    len(job_skills)
                    *
                    100
                )

            else:

                match_percentage = 0

            # Recommendations
            recommendations = generate_recommendations(
                missing_skills
            )

            # Match message
            if match_percentage >= 80:

                match_message = "Excellent Match!"

                match_description = (
                    "Your resume matches most of the "
                    "required skills for this job."
                )

            elif match_percentage >= 60:

                match_message = "Good Match!"

                match_description = (
                    "You have several relevant skills, "
                    "but there are some areas you can improve."
                )

            elif match_percentage >= 40:

                match_message = "Moderate Match"

                match_description = (
                    "You have some relevant skills, "
                    "but you should improve the missing skills."
                )

            else:

                match_message = "Needs Improvement"

                match_description = (
                    "Your resume has limited matching skills "
                    "for this job. Consider learning the missing skills."
                )

            # Create report filename
            report_filename = (
                "aws_job_analysis_report.pdf"
            )

            # Generate report
            generate_pdf_report(
                report_filename,
                match_percentage,
                matched_skills,
                missing_skills,
                recommendations
            )

            # Result
            result = {

                "match_percentage":
                    match_percentage,

                "matched_skills":
                    matched_skills,

                "missing_skills":
                    missing_skills,

                "recommendations":
                    recommendations,

                "match_message":
                    match_message,

                "match_description":
                    match_description,

                "report_filename":
                    report_filename
            }

        except Exception as e:

            error = (
                "Unable to analyze the resume. "
                "Please make sure the PDF or DOCX file is valid."
            )

            print("Error:", e)

    return render_template(
        "index.html",
        result=result,
        error=error
    )


# ============================================================
# Download Report
# ============================================================

@app.route("/download-report/<filename>")
def download_report(filename):

    return send_from_directory(
        app.config["REPORT_FOLDER"],
        filename,
        as_attachment=True
    )


# ============================================================
# Start Application
# ============================================================

if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )